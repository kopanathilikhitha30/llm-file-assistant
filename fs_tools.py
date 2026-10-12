"""File-system tools for reading, listing, writing, and searching resume files.

Supported input formats: PDF, TXT, and DOCX.
All public functions return JSON-serializable dictionaries/lists.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
import os

SUPPORTED_READ_EXTENSIONS = {".pdf", ".txt", ".docx"}


def _metadata(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "name": path.name,
        "filepath": str(path.resolve()),
        "extension": path.suffix.lower(),
        "size_bytes": stat.st_size,
        "modified_date": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
    }


def _extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts).strip()
    raise ValueError(f"Unsupported file extension: {ext or '(none)'}. Supported: PDF, TXT, DOCX.")


def read_file(filepath: str) -> dict[str, Any]:
    """Extract text and metadata from a PDF, TXT, or DOCX file."""
    path = Path(filepath).expanduser()
    try:
        if not path.exists():
            return {"success": False, "error": "File not found", "filepath": str(path)}
        if not path.is_file():
            return {"success": False, "error": "Path is not a file", "filepath": str(path)}
        if path.suffix.lower() not in SUPPORTED_READ_EXTENSIONS:
            return {"success": False, "error": f"Unsupported file type: {path.suffix}", "supported_extensions": sorted(SUPPORTED_READ_EXTENSIONS)}
        content = _extract_text(path)
        return {"success": True, "content": content, "metadata": _metadata(path), "character_count": len(content)}
    except Exception as exc:
        return {"success": False, "error": f"{type(exc).__name__}: {exc}", "filepath": str(path)}


def list_files(directory: str, extension: str | None = None) -> list[dict[str, Any]]:
    """List files directly inside a directory, optionally filtering by extension."""
    folder = Path(directory).expanduser()
    if not folder.exists():
        return [{"error": "Directory not found", "directory": str(folder)}]
    if not folder.is_dir():
        return [{"error": "Path is not a directory", "directory": str(folder)}]
    ext_filter = None
    if extension:
        ext_filter = extension.lower()
        if not ext_filter.startswith("."):
            ext_filter = "." + ext_filter
    result = []
    try:
        for path in sorted(folder.iterdir(), key=lambda p: p.name.lower()):
            if not path.is_file():
                continue
            if ext_filter and path.suffix.lower() != ext_filter:
                continue
            try:
                result.append(_metadata(path))
            except OSError as exc:
                result.append({"name": path.name, "error": str(exc)})
        return result
    except OSError as exc:
        return [{"error": f"{type(exc).__name__}: {exc}", "directory": str(folder)}]


def write_file(filepath: str, content: str) -> dict[str, Any]:
    """Write UTF-8 text, creating parent directories as needed."""
    path = Path(filepath).expanduser()
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return {"success": True, "message": "File written successfully", "metadata": _metadata(path)}
    except Exception as exc:
        return {"success": False, "error": f"{type(exc).__name__}: {exc}", "filepath": str(path)}


def search_in_file(filepath: str, keyword: str) -> dict[str, Any]:
    """Case-insensitively search extracted text and return matching snippets."""
    if not keyword or not keyword.strip():
        return {"success": False, "error": "Keyword must not be empty"}
    result = read_file(filepath)
    if not result.get("success"):
        return result
    content = result["content"]
    lower_content = content.lower()
    lower_keyword = keyword.lower()
    matches = []
    start = 0
    context_chars = 70
    while True:
        index = lower_content.find(lower_keyword, start)
        if index < 0:
            break
        left = max(0, index - context_chars)
        right = min(len(content), index + len(keyword) + context_chars)
        matches.append({
            "start": index,
            "end": index + len(keyword),
            "match": content[index:index + len(keyword)],
            "context": content[left:right].replace("\n", " "),
        })
        start = index + max(1, len(keyword))
    return {
        "success": True,
        "filepath": result["metadata"]["filepath"],
        "keyword": keyword,
        "match_count": len(matches),
        "matches": matches,
    }
